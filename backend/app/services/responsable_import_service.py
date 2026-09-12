from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.responsable import Responsable


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "nom_complet": (
        "nom et prenom",
        "nom & prenom",
        "nom prenom",
        "nom complet",
        "responsable",
        "directeur",
    ),
    "fonction": ("fonction", "poste", "role"),
    "grande_residence": (
        "grande residence",
        "grande direction",
        "residence",
        "structure",
    ),
}


@dataclass(slots=True)
class ResponsableImportRow:
    nom_complet: str
    fonction: str
    grande_residence: str
    source_file: str
    source_sheet: str
    source_row: int

    @property
    def normalized_name(self) -> str:
        return normalize_value(self.nom_complet)

    @property
    def normalized_fonction(self) -> str:
        return normalize_value(self.fonction)

    @property
    def normalized_grande_residence(self) -> str:
        return normalize_value(self.grande_residence)


@dataclass(slots=True)
class ResponsableImportSummary:
    parsed_count: int = 0
    inserted_count: int = 0
    updated_count: int = 0
    skipped_count: int = 0
    warnings: list[str] | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "parsed_count": self.parsed_count,
            "inserted_count": self.inserted_count,
            "updated_count": self.updated_count,
            "skipped_count": self.skipped_count,
            "warnings": self.warnings or [],
        }


class ResponsableImportService:
    def __init__(self, db: Session) -> None:
        self.db = db

    def parse_files(self, paths: list[Path]) -> tuple[list[ResponsableImportRow], list[str]]:
        rows: list[ResponsableImportRow] = []
        warnings: list[str] = []
        for path in paths:
            if path.suffix.lower() != ".xlsx":
                warnings.append(f"{path.name}: ignored because only .xlsx files are supported.")
                continue
            file_rows, file_warnings = self._parse_workbook(path)
            rows.extend(file_rows)
            warnings.extend(file_warnings)
        return rows, warnings

    def import_files(self, paths: list[Path]) -> ResponsableImportSummary:
        rows, warnings = self.parse_files(paths)
        summary = ResponsableImportSummary(parsed_count=len(rows), warnings=warnings)

        for row in rows:
            if not row.nom_complet or not row.fonction or not row.grande_residence:
                summary.skipped_count += 1
                continue

            existing = self.db.scalar(
                select(Responsable).where(
                    Responsable.normalized_name == row.normalized_name,
                    Responsable.normalized_fonction == row.normalized_fonction,
                    Responsable.normalized_grande_residence
                    == row.normalized_grande_residence,
                )
            )
            if existing is None:
                self.db.add(
                    Responsable(
                        nom_complet=row.nom_complet,
                        fonction=row.fonction,
                        grande_residence=row.grande_residence,
                        normalized_name=row.normalized_name,
                        normalized_fonction=row.normalized_fonction,
                        normalized_grande_residence=row.normalized_grande_residence,
                        source_file=row.source_file,
                        source_sheet=row.source_sheet,
                        source_row=row.source_row,
                    )
                )
                summary.inserted_count += 1
                continue

            existing.nom_complet = row.nom_complet
            existing.fonction = row.fonction
            existing.grande_residence = row.grande_residence
            existing.source_file = row.source_file
            existing.source_sheet = row.source_sheet
            existing.source_row = row.source_row
            summary.updated_count += 1

        self.db.commit()
        return summary

    def _parse_workbook(
        self,
        path: Path,
    ) -> tuple[list[ResponsableImportRow], list[str]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[ResponsableImportRow] = []
        warnings: list[str] = []
        for worksheet in workbook.worksheets:
            header_row, column_map = self._detect_header(worksheet)
            if header_row is None:
                warnings.append(
                    f"{path.name}/{worksheet.title}: no responsable header row was found."
                )
                continue
            max_column = max(column_map.values())
            for row_index, row in enumerate(
                worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_col=max_column,
                    values_only=True,
                ),
                start=header_row + 1,
            ):
                row_data = {
                    field_name: clean_text(
                        row[column_index - 1] if column_index - 1 < len(row) else None
                    )
                    for field_name, column_index in column_map.items()
                }
                if not any(row_data.values()):
                    continue
                rows.append(
                    ResponsableImportRow(
                        nom_complet=row_data.get("nom_complet", ""),
                        fonction=row_data.get("fonction", ""),
                        grande_residence=row_data.get("grande_residence", ""),
                        source_file=path.name,
                        source_sheet=worksheet.title,
                        source_row=row_index,
                    )
                )
        return rows, warnings

    def _detect_header(self, worksheet: Any) -> tuple[int | None, dict[str, int]]:
        best_row: int | None = None
        best_map: dict[str, int] = {}
        best_score = 0
        max_scan_row = min(20, worksheet.max_row or 20)
        for row_index in range(1, max_scan_row + 1):
            current_map: dict[str, int] = {}
            for cell in worksheet[row_index]:
                normalized = normalize_value(clean_text(cell.value))
                for field_name, aliases in HEADER_ALIASES.items():
                    if field_name in current_map:
                        continue
                    if any(alias in normalized for alias in aliases):
                        current_map[field_name] = cell.column
                        break
            score = len(current_map)
            if score > best_score:
                best_row = row_index
                best_map = current_map
                best_score = score
        if {"nom_complet", "fonction", "grande_residence"}.issubset(best_map):
            return best_row, best_map
        return None, {}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_value(value: str) -> str:
    text = unicodedata.normalize("NFKD", clean_text(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()
